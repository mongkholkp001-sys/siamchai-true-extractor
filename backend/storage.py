import os
import json
import re
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
RESULTS_DIR = os.path.join(BASE_DIR, "results")
CONFIG_FILE = os.path.join(DATA_DIR, "config.json")
USERS_FILE = os.path.join(DATA_DIR, "users.json")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(RESULTS_DIR, exist_ok=True)

DEFAULT_CONFIG = {
    # Web Dashboard Auth
    "web_username": "admin",
    "web_password": "password123",
    "auth_secret_token": "siamchai_true_secret_token_2026",
    
    # Siamchai Credentials
    "siamchai_username": "71481",
    "siamchai_password": "71482",
    
    # TrueCorp Credentials
    "true_username": "71114525",
    "true_password": "Slumzick999",
    
    # Crawler Settings
    "headless": True,
    "concurrency": True,
    "default_mode": "both", # "both", "siamchai", "true"
}

def load_config():
    cfg = DEFAULT_CONFIG.copy()
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                saved = json.load(f)
                cfg.update(saved)
        except Exception:
            pass

    # Read from environment variables if present (useful for Render / Docker / Cloud)
    for k in [
        "web_username", "web_password",
        "siamchai_username", "siamchai_password",
        "true_username", "true_password"
    ]:
        env_val = os.environ.get(k.upper()) or os.environ.get(k)
        if env_val:
            cfg[k] = env_val

    return cfg

def save_config(cfg):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=2, ensure_ascii=False)
        return True
    except Exception:
        return False

# ------------------ Multi-User Storage ------------------
def load_users():
    if os.path.exists(USERS_FILE):
        try:
            with open(USERS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list) and len(data) > 0:
                    return data
        except Exception:
            pass

    # If users.json doesn't exist, initialize with default admin
    cfg = load_config()
    default_admin = {
        "username": cfg.get("web_username", "admin").strip(),
        "password": cfg.get("web_password", "password123").strip(),
        "role": "admin",
        "name": "ผู้ดูแลระบบ (Admin)",
        "created_at": int(time.time())
    }
    save_users([default_admin])
    return [default_admin]

def save_users(users_list):
    try:
        with open(USERS_FILE, "w", encoding="utf-8") as f:
            json.dump(users_list, f, indent=2, ensure_ascii=False)
        return True
    except Exception:
        return False

def get_user(username: str):
    users = load_users()
    for u in users:
        if u.get("username", "").strip().lower() == username.strip().lower():
            return u
    return None

def add_user(username: str, password: str, role: str = "member", name: str = ""):
    users = load_users()
    clean_u = username.strip()
    if not clean_u:
        raise ValueError("ชื่อผู้ใช้ต้องไม่ว่างเปล่า")
    if get_user(clean_u):
        raise ValueError(f"ชื่อผู้ใช้ '{clean_u}' มีอยู่ในระบบแล้ว")
    
    new_user = {
        "username": clean_u,
        "password": password.strip(),
        "role": role if role in ["admin", "member"] else "member",
        "name": name.strip() or clean_u,
        "created_at": int(time.time())
    }
    users.append(new_user)
    save_users(users)
    return new_user

def update_user_password(username: str, new_password: str):
    users = load_users()
    clean_u = username.strip().lower()
    found = False
    for u in users:
        if u.get("username", "").strip().lower() == clean_u:
            u["password"] = new_password.strip()
            found = True
            break
    if not found:
        raise ValueError(f"ไม่พบผู้ใช้ '{username}'")
    save_users(users)
    return True

def delete_user(username: str, current_admin: str):
    clean_u = username.strip().lower()
    if clean_u == current_admin.strip().lower():
        raise ValueError("ไม่สามารถลบบัญชีตัวเองที่กำลังใช้งานอยู่ได้")
    users = load_users()
    admin_count = sum(1 for u in users if u.get("role") == "admin" and u.get("username", "").strip().lower() != clean_u)
    if admin_count < 1:
        raise ValueError("ไม่สามารถลบได้ เนื่องจากต้องมี Admin อย่างน้อย 1 บัญชีในระบบ")

    new_users = [u for u in users if u.get("username", "").strip().lower() != clean_u]
    if len(new_users) == len(users):
        raise ValueError(f"ไม่พบผู้ใช้ '{username}'")
    save_users(new_users)
    return True

def parse_cids_from_text(text: str):
    """Extract 13-digit Thai national IDs or numbers from text"""
    if not text:
        return []
    lines = text.strip().splitlines()
    cids = []
    for line in lines:
        cleaned = re.sub(r"[^\d]", "", line.strip())
        if cleaned:
            cids.append(cleaned)
    return cids

def parse_excel_cids(file_path: str):
    """Read first column from Excel or CSV file"""
    cids = []
    if not os.path.exists(file_path):
        return cids

    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".csv":
        try:
            with open(file_path, "r", encoding="utf-8-sig") as f:
                for line in f:
                    parts = line.strip().split(",")
                    if parts:
                        v = re.sub(r"[^\d]", "", parts[0].strip())
                        if v and len(v) >= 10:
                            cids.append(v)
            return cids
        except Exception:
            pass

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=1, max_col=1, values_only=True):
            val = row[0]
            if val is not None:
                cleaned = re.sub(r"[^\d]", "", str(val).strip())
                if cleaned and len(cleaned) >= 10:
                    cids.append(cleaned)
    except Exception as e:
        print(f"Error parsing Excel CIDs: {e}")

    return cids

def save_combined_results_to_excel(results, job_name="unified_extract"):
    """
    Save multi-platform results to a beautifully formatted Excel file with 3 sheets:
    1. สรุปรวม (Unified Summary)
    2. สยามชัย (Siamchai 17 Columns)
    3. ทรู (True Pre-Verify)
    """
    from datetime import datetime
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^\w\-]", "_", job_name)
    filename = f"{safe_name}_{timestamp}.xlsx"
    filepath = os.path.join(RESULTS_DIR, filename)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_fill_purple = PatternFill(start_color="6B21A8", end_color="6B21A8", fill_type="solid")
    header_fill_red = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Tahoma", size=9)
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # ---------------- SHEET 1: สรุปรวม (Summary) ----------------
    ws1 = wb.create_sheet(title="สรุปรวม")
    headers1 = [
        "ลำดับ", "เลขบัตรที่ใช้ค้นหา", "ชื่อ-นามสกุล", "สถานะสยามชัย",
        "เบอร์โทรผู้เช่าซื้อ (สยามชัย)", "จำนวนเบอร์ทรู (Active)", "รายการเบอร์ทรู (Active)",
        "เลขบัตรผู้ค้ำ", "ชื่อผู้ค้ำ", "เบอร์โทรผู้ค้ำ", "ที่อยู่ปัจจุบัน", "ที่ทำงาน"
    ]
    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(1, col_idx)
        cell.fill = header_fill_blue
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, r in enumerate(results, 1):
        cid = r.get("cid", "")
        sc = r.get("siamchai", {})
        tr = r.get("true", {})

        row_vals = [
            idx,
            f"'{cid}" if cid else "",
            sc.get("ชื่อ-นามสกุล") or tr.get("name", ""),
            sc.get("สถานะ", "-"),
            sc.get("เบอร์โทรผู้เช่าซื้อ", "-"),
            tr.get("active_count", 0),
            tr.get("active_phones", "-"),
            sc.get("เลขบัตรประชาชนผู้ค้ำ", "-"),
            sc.get("ชื่อ-นามสกุลผู้ค้ำ", "-"),
            sc.get("เบอร์โทรผู้ค้ำ", "-"),
            sc.get("ที่อยู่ปัจจุบันผู้เช่าซื้อ", "-"),
            sc.get("ที่อยู่ที่ทำงานผู้เช่าซื้อ", "-"),
        ]
        ws1.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws1.cell(idx + 1, c)
            cell.font = data_font
            cell.border = border_thin
            if c in [1, 2, 4, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # ---------------- SHEET 2: สยามชัย (17 คอลัมน์) ----------------
    ws2 = wb.create_sheet(title="สยามชัย")
    headers2 = [
        "ลำดับ", "เลขบัตรที่ใช้ค้นหา", "ชื่อ-นามสกุล", "สถานะ",
        "เลขบัตรประชาชนผู้เช่าซื้อ", "ชื่อ-นามสกุลผู้เช่าซื้อ", "เบอร์โทรผู้เช่าซื้อ",
        "ที่อยู่ตามบัตรประชาชนผู้เช่าซื้อ", "ที่อยู่ปัจจุบันผู้เช่าซื้อ",
        "ที่อยู่ที่ทำงานผู้เช่าซื้อ", "เบอร์โทรที่ทำงาน", "ตำแหน่ง",
        "เลขบัตรประชาชนผู้ค้ำ", "ชื่อ-นามสกุลผู้ค้ำ", "เบอร์โทรผู้ค้ำ",
        "ความสัมพันธ์", "ที่อยู่ผู้ค้ำ", "ที่ทำงานผู้ค้ำ"
    ]
    ws2.append(headers2)
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(1, col_idx)
        cell.fill = header_fill_purple
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, r in enumerate(results, 1):
        sc = r.get("siamchai", {})
        cid = r.get("cid", "")
        row_vals = [
            idx,
            f"'{cid}" if cid else "",
            sc.get("ชื่อ-นามสกุล", "-"),
            sc.get("สถานะ", "-"),
            sc.get("เลขบัตรประชาชนผู้เช่าซื้อ", "-"),
            sc.get("ชื่อ-นามสกุลผู้เช่าซื้อ", "-"),
            sc.get("เบอร์โทรผู้เช่าซื้อ", "-"),
            sc.get("ที่อยู่ตามบัตรประชาชนผู้เช่าซื้อ", "-"),
            sc.get("ที่อยู่ปัจจุบันผู้เช่าซื้อ", "-"),
            sc.get("ที่อยู่ที่ทำงานผู้เช่าซื้อ", "-"),
            sc.get("เบอร์โทรที่ทำงาน", "-"),
            sc.get("ตำแหน่ง", "-"),
            sc.get("เลขบัตรประชาชนผู้ค้ำ", "-"),
            sc.get("ชื่อ-นามสกุลผู้ค้ำ", "-"),
            sc.get("เบอร์โทรผู้ค้ำ", "-"),
            sc.get("ความสัมพันธ์", "-"),
            sc.get("ที่อยู่ผู้ค้ำ", "-"),
            sc.get("ที่ทำงานผู้ค้ำ", "-"),
        ]
        ws2.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws2.cell(idx + 1, c)
            cell.font = data_font
            cell.border = border_thin
            if c in [1, 2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # ---------------- SHEET 3: ทรู (True Pre-Verify) ----------------
    ws3 = wb.create_sheet(title="ทรู")
    headers3 = [
        "ลำดับ", "เลขบัตรที่ใช้ค้นหา", "สถานะ", "จำนวนเบอร์ Active",
        "รายการเบอร์โทรศัพท์ (Active)", "เบอร์โทรศัพท์ทั้งหมดที่พบ", "สถานะทุกเบอร์"
    ]
    ws3.append(headers3)
    for col_idx in range(1, len(headers3) + 1):
        cell = ws3.cell(1, col_idx)
        cell.fill = header_fill_red
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, r in enumerate(results, 1):
        tr = r.get("true", {})
        cid = r.get("cid", "")
        row_vals = [
            idx,
            f"'{cid}" if cid else "",
            tr.get("status", "-"),
            tr.get("active_count", 0),
            tr.get("active_phones", "-"),
            tr.get("all_phones", "-"),
            tr.get("phone_details", "-"),
        ]
        ws3.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws3.cell(idx + 1, c)
            cell.font = data_font
            cell.border = border_thin
            if c in [1, 2, 3, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto column width for all sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    wb.save(filepath)
    return filepath

def save_siamchai_results_to_excel(results, job_name="siamchai_extract"):
    """
    Save Siamchai 17-column results to a dedicated Excel file.
    """
    from datetime import datetime
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^\w\-]", "_", job_name)
    filename = f"{safe_name}_{timestamp}.xlsx"
    filepath = os.path.join(RESULTS_DIR, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "สยามชัย (17 คอลัมน์)"

    header_fill = PatternFill(start_color="6B21A8", end_color="6B21A8", fill_type="solid")
    header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Tahoma", size=9)
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    headers = [
        "ลำดับ", "เลขบัตรที่ใช้ค้นหา", "ชื่อ-นามสกุล", "สถานะ",
        "เลขบัตรประชาชนผู้เช่าซื้อ", "ชื่อ-นามสกุลผู้เช่าซื้อ", "เบอร์โทรผู้เช่าซื้อ",
        "ที่อยู่ตามบัตรประชาชนผู้เช่าซื้อ", "ที่อยู่ปัจจุบันผู้เช่าซื้อ",
        "ที่อยู่ที่ทำงานผู้เช่าซื้อ", "เบอร์โทรที่ทำงาน", "ตำแหน่ง",
        "เลขบัตรประชาชนผู้ค้ำ", "ชื่อ-นามสกุลผู้ค้ำ", "เบอร์โทรผู้ค้ำ",
        "ความสัมพันธ์", "ที่อยู่ผู้ค้ำ", "ที่ทำงานผู้ค้ำ"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, r in enumerate(results, 1):
        sc = r.get("siamchai", {})
        cid = r.get("cid", "")
        row_vals = [
            idx,
            f"'{cid}" if cid else "",
            sc.get("ชื่อ-นามสกุล", "-"),
            sc.get("สถานะ", "-"),
            sc.get("เลขบัตรประชาชนผู้เช่าซื้อ", "-"),
            sc.get("ชื่อ-นามสกุลผู้เช่าซื้อ", "-"),
            sc.get("เบอร์โทรผู้เช่าซื้อ", "-"),
            sc.get("ที่อยู่ตามบัตรประชาชนผู้เช่าซื้อ", "-"),
            sc.get("ที่อยู่ปัจจุบันผู้เช่าซื้อ", "-"),
            sc.get("ที่อยู่ที่ทำงานผู้เช่าซื้อ", "-"),
            sc.get("เบอร์โทรที่ทำงาน", "-"),
            sc.get("ตำแหน่ง", "-"),
            sc.get("เลขบัตรประชาชนผู้ค้ำ", "-"),
            sc.get("ชื่อ-นามสกุลผู้ค้ำ", "-"),
            sc.get("เบอร์โทรผู้ค้ำ", "-"),
            sc.get("ความสัมพันธ์", "-"),
            sc.get("ที่อยู่ผู้ค้ำ", "-"),
            sc.get("ที่ทำงานผู้ค้ำ", "-"),
        ]
        ws.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws.cell(idx + 1, c)
            cell.font = data_font
            cell.border = border_thin
            if c in [1, 2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    wb.save(filepath)
    return filepath

def save_true_results_to_excel(results, job_name="true_extract"):
    """
    Save TrueCorp Active Number results to a dedicated Excel file.
    """
    from datetime import datetime
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^\w\-]", "_", job_name)
    filename = f"{safe_name}_{timestamp}.xlsx"
    filepath = os.path.join(RESULTS_DIR, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ทรู (Active Numbers)"

    header_fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Tahoma", size=9)
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    headers = [
        "ลำดับ", "เลขบัตรที่ใช้ค้นหา", "สถานะ", "จำนวนเบอร์ Active",
        "รายการเบอร์โทรศัพท์ (Active)", "เบอร์โทรศัพท์ทั้งหมดที่พบ", "สถานะทุกเบอร์"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, r in enumerate(results, 1):
        tr = r.get("true", {})
        cid = r.get("cid", "")
        row_vals = [
            idx,
            f"'{cid}" if cid else "",
            tr.get("status", "-"),
            tr.get("active_count", 0),
            tr.get("active_phones", "-"),
            tr.get("all_phones", "-"),
            tr.get("phone_details", "-"),
        ]
        ws.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws.cell(idx + 1, c)
            cell.font = data_font
            cell.border = border_thin
            if c in [1, 2, 3, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    wb.save(filepath)
    return filepath

